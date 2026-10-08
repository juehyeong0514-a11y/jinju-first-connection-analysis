"""Fill the structure of official attachment 2, keeping the v5 calculations.

Input DOCX was converted from the unmodified official HWP via pyhwp/LibreOffice.
Its lost logo/font/column mappings are restored from the HWP preview. Render the
output with render_docx.py, then use that exact PDF as the submission version.
"""
import copy
import json
import os
import subprocess
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.table import Table
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT = Path(__file__).resolve().parents[1]
TMP = ROOT / 'tmp/official'; TMP.mkdir(parents=True, exist_ok=True)
D = json.loads((ROOT / 'outputs/robust/analysis.json').read_text())
E = json.loads((ROOT / 'outputs/decisions/analysis.json').read_text())
IDENTITY = json.loads((ROOT / 'submission/identity.json').read_text())
CATALOG = json.loads((ROOT / 'submission/data_catalog.json').read_text())
SOURCE = ROOT / 'data/raw/contest_report_template.docx'
doc = Document(SOURCE)
if 'Title' not in doc.styles: doc.styles.add_style('Title', WD_STYLE_TYPE.PARAGRAPH)
original_rows = [copy.deepcopy(r._tr) for r in doc.tables[1].rows]
table_props = copy.deepcopy(doc.tables[1]._tbl.tblPr)
# The converted two-column grid was equal-width; the HWP preview is 22/128mm.
body = doc._element.body
for child in list(body):
    if child.tag != qn('w:sectPr'): body.remove(child)
FONT = os.environ.get('KR_DOC_FONT', 'AppleGothic')

def xml(parent, tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items(): e.set(qn(k), str(v))
    parent.append(e)
    return e

def fmt(p, size=10, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, after=3):
    p.alignment = align
    f = p.paragraph_format
    f.space_before = Pt(0); f.space_after = Pt(after)
    f.line_spacing = 1.25; f.keep_with_next = False; f.widow_control = True
    f.left_indent = Pt(0); f.right_indent = Pt(0); f.first_line_indent = Pt(0)
    for r in p.runs:
        r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold
        r.font.italic = False; r.font.color.rgb = RGBColor(0, 0, 0)
        rf = r._element.get_or_add_rPr().get_or_add_rFonts()
        for key in ['w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs']: rf.set(qn(key), FONT)
    return p

def para(cell, text='', size=10, bold=False):
    p = cell.paragraphs[0] if len(cell.paragraphs)==1 and not cell.paragraphs[0].text else cell.add_paragraph()
    p.add_run(text); return fmt(p, size, bold)

def hyperlink(p, url, size=8.5):
    relationship=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
    h=xml(p._p,'w:hyperlink',**{'r:id':relationship})
    r=xml(h,'w:r');pr=xml(r,'w:rPr')
    rf=xml(pr,'w:rFonts')
    for key in ['w:ascii','w:hAnsi','w:eastAsia','w:cs']:rf.set(qn(key),FONT)
    xml(pr,'w:sz',**{'w:val':int(size*2)});xml(pr,'w:color',**{'w:val':'000000'})
    xml(r,'w:t').text=url

def clean(cell):
    for child in list(cell._tc):
        if child.tag != qn('w:tcPr'): cell._tc.remove(child)
    cell._tc.append(OxmlElement('w:p'))
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

def table_page():
    t = doc.add_table(rows=0, cols=2); t.autofit = False
    t._tbl.remove(t._tbl.tblPr); t._tbl.insert(0, copy.deepcopy(table_props))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for c, width in zip(t.columns, [2.2, 12.8]): c.width = Cm(width)
    return t

def slot(t, index, label=None):
    rowxml = copy.deepcopy(original_rows[index])
    t._tbl.append(rowxml); row = t.rows[-1]
    for c, width in zip(row.cells, [2.2, 12.8]):
        c.width = Cm(width); clean(c)
        pr = c._tc.get_or_add_tcPr()
        mar = pr.find(qn('w:tcMar'))
        if mar is not None: pr.remove(mar)
        mar = xml(pr, 'w:tcMar')
        for edge in ['top','bottom']: xml(mar, 'w:'+edge, **{'w:w': '75', 'w:type':'dxa'})
        for edge in ['start','end']: xml(mar, 'w:'+edge, **{'w:w':'85', 'w:type':'dxa'})
    para(row.cells[0], label or ['이름/팀명','주제','분석 배경\n및 목적','활용\n데이터','분석과정\n및 방법','AI 기술\n활용','분석 내용\n및 결과','정책제안\n및 기대효과','분석도구\n및 참고문헌','분석결과\n이미지\n(필수)'][index], 10)
    row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return row.cells[1]

def compact_table(cell, headers, rows, widths, size=8.5):
    t = cell.add_table(rows=1, cols=len(headers)); t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for col,w in zip(t.columns,widths): col.width = Cm(w)
    for i, values in enumerate([headers]+rows):
        cells = t.rows[0].cells if i==0 else t.add_row().cells
        for c,w,text in zip(cells,widths,values):
            c.width=Cm(w); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=c.paragraphs[0];p.add_run(str(text));fmt(p,size,i==0,after=1)
            pr=c._tc.get_or_add_tcPr(); borders=xml(pr,'w:tcBorders')
            for edge in ['top','bottom','start','end']:
                xml(borders,'w:'+edge,**{'w:val':'single','w:sz':'3','w:color':'777777'})
            mar=xml(pr,'w:tcMar')
            for edge in ['top','bottom']:xml(mar,'w:'+edge,**{'w:w':'35','w:type':'dxa'})
            for edge in ['start','end']:xml(mar,'w:'+edge,**{'w:w':'45','w:type':'dxa'})
            if i==0:xml(pr,'w:shd',**{'w:fill':'F2F2F2'})
        if i==0:xml(t.rows[0]._tr.get_or_add_trPr(),'w:tblHeader')
        xml(t.rows[i]._tr.get_or_add_trPr(),'w:cantSplit')
    for p in cell.paragraphs:
        if not p.text: fmt(p,1,after=0)
    return t

def page_break():
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0)
    p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1
    p.paragraph_format.page_break_before=True
    p.paragraph_format.keep_with_next=True
    p.add_run().font.size=Pt(1)

# Restore the organizer banner from the original HWP, not replacement branding.
logo=doc.add_table(rows=1, cols=2);logo.autofit=False
for c in logo.rows[0].cells:c.width=Cm(7.5)
for i,(name,width) in enumerate([('BIN0002.png',2.4),('BIN0001.png',4.2)]):
    p=logo.cell(0,i).paragraphs[0];fmt(p,after=3)
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 else WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run().add_picture(str(ROOT/'data/raw/contest_report_assets'/name),width=Cm(width))
p=doc.add_paragraph('- AI와 함께하는 교통문제 해결을 위한 데이터 분석 공모전 -')
fmt(p,10.5,align=WD_ALIGN_PARAGRAPH.CENTER,after=1)
ppr=p._p.get_or_add_pPr();borders=xml(ppr,'w:pBdr');xml(borders,'w:top',**{'w:val':'single','w:sz':'24','w:space':'5','w:color':'85B67B'})
p=doc.add_paragraph('분석보고서');p.style=doc.styles['Title']
fmt(p,14,True,WD_ALIGN_PARAGRAPH.CENTER,8)
ppr=p._p.get_or_add_pPr();borders=xml(ppr,'w:pBdr');xml(borders,'w:bottom',**{'w:val':'single','w:sz':'24','w:space':'5','w:color':'85B67B'})

# Page 1: original identity/background/data slots. All source URLs are here.
t=table_page();c=slot(t,0);para(c,IDENTITY['participant']+' / 개인 참가')
c=slot(t,1);para(c,'진주 주거지의 서울 조기 도착 접근성 분석과 장거리 첫차 연계 개선',bold=True)
c=slot(t,2)
para(c,'장거리 첫차가 있어도 승차장 접근 보행과 시내버스 첫 운행이 맞지 않으면 이른 서울 도착에 제약이 생긴다. 진주 5개 생활권의 대표 정류장 20곳에서 강남역·서울역·사당역을 비교해 개별 보행 부담에 맞는 접근 지원과 운행 조정의 조건을 제시하고자 했다.',9.5)
para(c,'핵심 두 지점의 평일 강남 도착은 보행 20분 이내 조건에서 11:00→09:27이다. 40분 보행을 허용하면 기존 첫차에 연결돼 추가 개선은 0분이다. 지원이 필요한 조건을 구분한다.',9.5)
c=slot(t,3)
para(c,'공개자료 조회일 2026.10.08. 번호는 본문 인용이며 지도 보행·예상 운임은 참고값이다.',8.5)
rows=[]
for r in CATALOG:
    if r['reference']=='15':continue
    provider=r['provider'].replace('경상남도 진주시','진주시').replace('전국고속버스운송사업조합','전국고속버스\n운송사업조합')
    rows.append(['['+r['reference']+'] '+r['data_name'],provider,r['platform'],r['url'],r['selection_reason']])
compact_table(c,['활용 데이터(명)','제공기관(명)','출처\n플랫폼(명)','URL','선정 이유'],rows,[2.85,1.45,1.65,4.9,1.65],7.4)

# Page 2: methodology and concrete AI use, as separate official fields.
page_break();t=table_page();c=slot(t,4)
para(c,'1. 수집과 범위',bold=True)
para(c,'평일 10/8·13·15, 토요일 10/17, 일요일 10/18의 버스·철도와 서울 전철 시각표를 결합했다. 생활권별 4곳을 사전에 선정한 20개 표본이며 인구 대표 표본은 아니다. 고속·시외버스 API의 미래 날짜 빈 응답은 무운행으로 간주하지 않고 KOBUS 공개 배차와 공식 열차표로 대조했다. 미래 일반열차와 임시 변경은 불완전하다. [1-5,9,10]')
para(c,'2. 전처리와 결측 관리',bold=True)
para(c,'정류장 ID·좌표·순서·방향·요일을 대응했다. 노선번호 묶음의 가장 이른 출발과 지선 요약 첫차를 잘못 비교해 제외했던 50방향을 복원했다. 미대응 추천 471개 중 127개는 좌표 3m 이내, 앞 정류장 순서 및 방향 유일 조건으로 복원했다. 직접 후보 848개와 대응 추천 1,926개를 검토하되 미대응 344개(핵심 22개)를 남겼다. 결측을 임의로 채우지 않았다. [1,6]')
para(c,'3. 시간순 경로 탐색과 비교 지표',bold=True)
para(c,'04:00 이후 출발, 지역 외부 접근 보행 합계 20분 이내, 지역 승차 여유 2분을 기본 조건으로 삼았다. 허브 도착에 승강장 이동 5분과 장거리 승차 여유 10분을 더해 탈 수 있는 버스·열차를 선택했다. 서울 전철은 같은 열차번호의 승하차를 연결하고 명시한 환승·지상 보행을 더했다. 알려진 후보와 서울 회랑에서 최종 도착이 가장 이른 경로를 비교한 것으로, 전체 교통망의 완전 탐색 최적해는 아니다. [1-6]')
para(c,'주 지표는 기존 도착시각−개선 도착시각(분), 목표시각까지 가능한 표본 수, 출발시각 및 여정 길이이다. 접근성 지표에 따라 목적지 기회를 비교했다. [15] 첫차 50분 전에 출발하는 연계편, 기존 노선 조정, 철도 접근, 직접 보행, 예약택시를 비교하고 세 지선 배정 조건을 사용했다. 첫·막차만 사용하는 보수적 경계, 같은 첫차·비고 후보를 지선 운행 범위에 대응한 중간 조건, 전체 묶음표를 허용하는 낙관적 경계다. 후속 지선 운행의 실제 배정은 확인되지 않았다.')
para(c,'4. 민감도와 검증',bold=True)
para(c,'요일·목적지·보행 한계 10/20/30/40분·최초 출발 가능시각·목표시각·도로 도착 지연·전철 진입 보행을 바꾸었다. 지연마다 기존안과 개선안 모두 버스·철도 및 예정 전철을 다시 선택했다. 운영비는 차량 좌석, 가족·기존 일행, 사전 취소와 늦은 불참을 반영했다. 내부 논리 검사 36,105개 및 추가 검사 12,867개를 통과했지만 실제 수요·정시율·견적의 검증을 뜻하지 않는다.')
c=slot(t,5)
para(c,'사용 서비스는 OpenAI Codex이다. 공개자료 수집, 정류장 대응, 시간표 연결, 민감도와 반례 점검, Python 코드 및 보고서 작성에 활용했다. AI가 제안한 해석은 원문과 계산 결과를 대조했고, 예외 경로와 보행 대안에 따라 결론을 수정했다.')
para(c,'주요 프롬프트: “어느 주거지에서 서울의 같은 목적지에 더 일찍 도착할 수 있고, 첫 운행을 어떻게 조정하면 그 차이가 줄어드는가?”, “버스 외 철도 대안을 포함”, “공개자료 중심으로 규모·누락·요일/목적지·운영비를 보강”. 주말 고정안 실패, 긴 여정, 늦은 출발, 지연, 동승·좌석, 직접 보행과 동일 첫차 택시를 반례로 점검했다. 참가자가 자료와 해석을 최종 확인하며, 결과 재계산에는 AI 접속이 필요하지 않다.')

# Page 3: results, including counterexamples that condition the finding.
page_break();t=table_page();c=slot(t,6)
para(c,'1. 같은 목적지의 이른 도착 기회',bold=True)
compact_table(c,['요일 / 핵심 두 지점','강남역','서울역','사당역'],[
 ['평일 10/8·13·15','11:00→09:27\n93분 앞당김','11:21→09:44\n96.5분 앞당김','11:04→09:33\n91분 앞당김'],
 ['토 10/17 지선 배정 경계','82~102분\n연계 09:37','78~102분\n연계 09:58','82~99분\n연계 09:44'],
 ['일 10/18','10:59→09:05\n114분 앞당김','11:16→09:28\n108분 앞당김','11:06→09:11\n115.5분 앞당김']], [3.5,3,3,3],8.5)
para(c,'LH10단지(49008)·풀에버정문(49036), 지역 보행 20분 이하 조건이다. 시각은 분 올림, 차이는 초까지 계산했다. 토요일은 실제 지선 배정의 불확실성 경계다. 300개 비교 중 36개가 배정 가정의 영향을 받았다. 10/17 첫차 잔여 9석은 조회 순간 정보이며 연계편 12석이 장거리 12석을 보장하지 않는다. [1-5,9]',8.5)
para(c,'2. 보행 한계와 출발 부담',bold=True)
wr=[r for r in E['walking_sensitivity'] if r['date']=='2026-10-08' and r['target']=='강남' and r['branch_mode']=='matched']
rows=[]
for cap in [10,20,30,40]:
    a=next(r for r in wr if r['walk_cap']==cap and r['origin_id']=='49008');b=next(r for r in wr if r['walk_cap']==cap and r['origin_id']=='49036')
    rows.append([str(cap)+'분',f"{a['advance']:g}분",f"{b['advance']:g}분",f"{a['baseline_count_0930']}→{a['feeder_count_0930']}곳"])
compact_table(c,['허용 지역 보행','LH10 앞당김','풀에버 앞당김','09:30 기존→연계'],rows,[3.2,2.8,2.8,3.7])
para(c,'20분은 조사된 주민 선호나 공식 지원 기준이 아니다. 대표 정류장 보행은 26/32분, 별도 상세 지도 출발점은 28/31분이며 서로 대체할 수 없다. 40분을 허용하면 두 지점 모두 기존 첫차로 09:27에 도착하여 추가 개선이 0분이다. 야간 보행 환경과 개인별 보행 능력은 실측하지 않았다. [6]',9)
para(c,'같은 상세 지도 지점의 평일 확인 직통 후보는 LH10 06:23→11:00(277분), 풀에버 06:21→11:00(279분)이다. 연계안은 각각 04:38/04:36→09:27(289/291분)으로 105분 일찍 출발하고 여정은 12분 늘어난다. 05:00부터 출발할 수 있는 이용자는 연계편을 놓쳐 11:00 도착에 그친다. 전체 경로의 가장 늦은 출발 최적해를 의미하지 않는다.',9)
para(c,'3. 목표시각과 지연의 영향',bold=True)
compact_table(c,['목표시각','09:00','09:30','10:00','10:30','11:00'],[['기존→연계','4→4곳','6→8곳','6→8곳','10→12곳','19→19곳']],[2.5,2,2,2,2,2],8.5)
para(c,'평일 강남 20곳, 지역 보행 20분 조건이다. 목표시각은 비교 격자이며 관측한 근로·예약 수요가 아니다. 09:30 목표의 기본 여유는 3분이다. 모든 장거리버스 서울 도착에 +10분을 적용하면 핵심 연계 도착은 09:37, +30분은 09:55, +60분은 10:31이며 기존안은 철도로 바뀌어 개선 폭은 54.5분이다. 전철 진입 보행 +10분도 09:37이다. 가상 지연 조건이며 발생확률·정시율을 추정하지 않았다.',9)
para(c,'4. 지역 규모와 수요의 구분',bold=True)
para(c,'관련 두 단지는 한림풀에버 1,421호와 NHF10 404호, 합계 1,825호다. 충무공동 인구 33,698명·12,904세대는 행정동 전체 규모이다. 모두 실제 수혜 가구·탑승 수요로 환산하지 않았다. ITS 9월 150번 승차 68,904건 중 05시대 176건(하루 5.9건)은 기존 노선 전체의 승차로 신규 서울행 수요가 아니다. 시간·일·노선 합계 1,983,808건과 행정동 합계의 차이 3,697건은 유지했다. [7,12,13]',9)

# Page 4: policy, effects, tools and references in the original order.
page_break();t=table_page();c=slot(t,7)
para(c,'1. 개별 보행 부담을 확인한 첫차 예약 접근 지원',bold=True)
para(c,'주소만으로 지원 필요를 판단하지 않는다. 개별 보행 가능성과 부담, 출발시각을 확인하고 장거리 승차권을 먼저 확보한다. 확정 일행·좌석·취소 조건과 회송·인건비·보험을 포함한 동일 범위 견적에서 예약택시 지원 또는 연계편을 선택한다. 기대효과는 보행 부담을 제한할 때의 이른 도착 선택지 확대다. 시간 절약액·사업 편익비·탄소감축량으로 환산하지 않는다.',9)
compact_table(c,['요일','혁신 첫 서울행','풀에버→LH10→승차장'],[
 ['평일','05:30','04:40→04:46→04:55'],['토요일','05:40','04:50→04:56→05:05'],['일요일','05:10','04:20→04:26→04:35']],[2.4,3,7.1])
para(c,'연계편은 첫 서울행 50분 전 출발, 주행 15분 계획이다. 주행 22.5분·승강장 이동 10분·승차 여유 15분 조건의 잔여는 2.5분이다. 실측 성공률은 아니며 기존 04:45 고정안은 일요일 첫차를 놓친다. 하모콜버스 공식 안내의 외곽형 동부 5면 06~22시·관광형 09~22시는 이 첫차 시간에 앞선 연결을 제공하지 않는다. 기존 차량의 새벽 사용권이나 운영 승인은 미확인이다. [9,14]',9)
para(c,'2. 예약택시와 차량 규모를 함께 비교',bold=True)
para(c,'지도 지점에 택시 확보·대기 0분·주행 5분·승강장 이동 5분·승차 여유 10분이면 05:10 출발로 같은 첫차를 타고 09:27 도착한다. 연계편보다 32~34분 늦게 출발한다. 직접 보행 28/31분도 같은 첫차에 연결된다. 택시 지원은 접근 비용과 보행 부담을 바꾼다. 새벽 차량·예약료·견적은 미검증이며 지도 4,600/5,000원은 계약 가격이 아니다. [6]',9)
compact_table(c,['12명 조건 / 취소·불참 없음','연계 총비용','택시 총비용'],[
 ['독립 12일행 / 연계 12석','6만원','6만원'],['2명씩 가족·기존 6일행 / 연계 12석','6만원','3만원'],['독립 12일행 / 연계 8석·2대','12만원','6만원']],[7.1,2.7,2.7])
para(c,'가정: 연계 6만원/대, 택시 5천원/대·3석, 본인부담 1,650원/실제 탑승자. 모르는 일행을 임의 동승시키지 않는다. 차량 수는 확정 인원/좌석의 올림이며 늦은 불참으로 계약 차량비는 줄지 않는다. 지원비=max(0, 총비용−실제 탑승자 본인부담). 차량 3/6/9/12만원, 택시 4,600/5,000/8,000원, 연계 8/12/16석, 일행 1/2/3명 및 취소·불참의 8,100개 조건을 계산했다. 실제 견적은 별도 확인한다.',9)
para(c,'3. 4주 시범운영과 확대 판단',bold=True)
para(c,'4주 운영은 제안이며 수행하지 않았다. 보행 부담·예약불가·장거리 좌석, 취소·불참·차량 준비·승차장 도착·장거리 승차·서울 도착·원가를 기록한다. 같은 서비스 범위의 비용과 실제 목표 충족을 확인한 뒤 확대 여부를 결정한다. 수요와 야간 보행 제약 미확인 단계에서는 포괄 지원을 주장하지 않는다.',9)
c=slot(t,8)
para(c,'분석도구: Python 3(CSV·JSON 처리와 시간순 경로 탐색), openpyxl(철도표 전처리), ReportLab(그림·검증 부록), python-docx(양식 작성). 원자료 목록과 출처는 1쪽 [1-10,12-14] 표에 제시했다.',9)
para(c,'[15] OECD/ITF(2019), Improving Transport Planning and Investment Through the Use of Accessibility Indicators. 접근성 방법론 참고이며 진주 실증 근거는 아니다. https://itf-oecd.org/transport-planning-investment-accessibility-indicators',8.5)
p=para(c,'공개 분석 코드: ',8.5);hyperlink(p,IDENTITY['code_url'])
para(c,'공개 저장소는 코드와 집계표를 제공한다. 전체 원문 경로 재현에는 별도 로컬 캐시 묶음과 RUN_ANALYSIS.txt가 필요하다. 인증키는 제외했다. 재현 검사는 실제 운행 성공을 입증하지 않는다.',8.5)

# Page 5: mandatory image field, drawn directly from calculated results.
page_break();t=table_page();c=slot(t,9)
FONT_PATH=Path(os.environ.get('KR_FONT','/System/Library/Fonts/Supplemental/AppleGothic.ttf'))
pdfmetrics.registerFont(TTFont('OfficialKR',str(FONT_PATH)))
def scn(policy):
    return next(e for e in D['experiments'] if all(e['parameters'][k]==v for k,v in [('date','2026-10-08'),('target','강남'),('branch_mode','matched'),('policy',policy)]))
B=scn('baseline');F=scn('feeder')
def make_plot(name,w,h,draw):
    pdf=TMP/(name+'.pdf');cc=canvas.Canvas(str(pdf),pagesize=(w,h));draw(cc);cc.save()
    subprocess.run(['pdftoppm','-singlefile','-r','180','-png',str(pdf),str(TMP/name)],check=True,capture_output=True)
    return TMP/(name+'.png')
def arrival(cc):
    left=132;right=346;bottom=24;top=324
    x=lambda v:left+(v-510)/170*(right-left)
    cc.setFont('OfficialKR',8)
    for v in [510,540,570,600,630,660]:
        cc.setStrokeColorRGB(.83,.83,.83);cc.setLineWidth(.4);cc.line(x(v),bottom,x(v),top)
        cc.setFillColorRGB(0,0,0);cc.drawCentredString(x(v),10,f'{int(v)//60:02d}:{int(v)%60:02d}')
    cc.setStrokeColorRGB(.7,.35,.15);cc.setDash(3,2);cc.line(x(570),bottom,x(570),top);cc.setDash()
    for i,(a,b) in enumerate(zip(B['results'],F['results'])):
        y=top-10-i*14.5;cc.setFillColorRGB(0,0,0);cc.setFont('OfficialKR',7.8)
        cc.drawString(2,y-3,a['zone']+'  '+a['origin'])
        cc.setStrokeColorRGB(.0,.48,.46);cc.setLineWidth(1.6);cc.line(x(a['arrival']),y,x(b['arrival']),y)
        cc.setFillColorRGB(.12,.22,.3);cc.circle(x(a['arrival']),y,2.7,fill=1,stroke=0)
        if a['arrival']!=b['arrival']:
            cc.setFillColorRGB(.0,.48,.46);cc.circle(x(b['arrival']),y,2.7,fill=1,stroke=0)
            cc.setFont('OfficialKR',7);cc.drawString(x(b['arrival'])+5,y+3,'93분')
def walk(cc):
    left=42;right=334;bottom=24;top=121
    xx=lambda cap:left+(cap-10)/30*(right-left)
    yy=lambda val:bottom+val/100*(top-bottom)
    cc.setFont('OfficialKR',8)
    for v in [0,30,60,90]:
        cc.setStrokeColorRGB(.83,.83,.83);cc.setLineWidth(.4);cc.line(left,yy(v),right,yy(v));cc.setFillColorRGB(0,0,0);cc.drawRightString(left-5,yy(v)-3,str(v))
    for cap in [10,20,30,40]:cc.drawCentredString(xx(cap),9,str(cap)+'분')
    for oid,color,label,y in [('49008',(.12,.22,.3),'LH10단지',143),('49036',(.0,.48,.46),'풀에버정문',132)]:
        vals=[next(r['advance'] for r in wr if r['walk_cap']==cap and r['origin_id']==oid) for cap in [10,20,30,40]]
        cc.setStrokeColorRGB(*color);cc.setFillColorRGB(*color);cc.setLineWidth(1.5)
        for i in range(3):cc.line(xx([10,20,30][i]),yy(vals[i]),xx([20,30,40][i]),yy(vals[i+1]))
        for cap,val in zip([10,20,30,40],vals):cc.circle(xx(cap),yy(val),2.5,fill=1,stroke=0)
        cc.drawString(220,y,label)
    cc.setFillColorRGB(0,0,0);cc.drawString(2,144,'도착 앞당김(분)')
para(c,'그림 1 진주 20개 대표 정류장의 강남 도착시각 비교',10,True)
img=make_plot('arrival',355,336,arrival);p=c.add_paragraph();fmt(p,after=2);p.add_run().add_picture(str(img),width=Cm(12.2))
para(c,'10/8 평일, 지역 보행 20분 이하, 세 지선 조건 중 중간 배정. 남색은 기존, 청록은 첫차 연계편 추가, 점선은 09:30이다. 핵심 두 지점은 11:00→09:27, 09:30 도착 가능 표본은 6→8곳이다. 경로 후보의 시간표 계산이며 실제 도착 기록은 아니다. [1-6,9]',9)
para(c,'그림 2 보행 허용 조건에 따른 핵심 두 지점의 추가 도착 개선',10,True)
img=make_plot('walking',355,157,walk);p=c.add_paragraph();fmt(p,after=2);p.add_run().add_picture(str(img),width=Cm(12.2))
para(c,'10/8 평일 강남 비교. 30분 보행에서는 LH10의 개선이 0분, 40분에서는 두 지점 모두 0분이다. “93분 개선”은 지역 보행 20분 이하 가정의 결과이다. 주민의 실제 보행 능력과 선호를 측정한 기준이 아니므로, 예약 지원 전에 개별 보행 대안과 부담을 확인해야 한다. [6]',9)

# Restore safe font mappings throughout converted styles; keep original geometry.
for style in doc.styles:
    if style.type==1:
        style.font.name=FONT;style.font.color.rgb=RGBColor(0,0,0)
        rf=style._element.get_or_add_rPr().get_or_add_rFonts()
        for key in ['w:ascii','w:hAnsi','w:eastAsia','w:cs']:rf.set(qn(key),FONT)
for section in doc.sections:
    section.footer_distance=Cm(.7)
    p=section.footer.paragraphs[0];p.clear();p.add_run('- ')
    field=xml(p._p,'w:fldSimple',**{'w:instr':'PAGE'});r=xml(field,'w:r');xml(r,'w:t').text='1'
    p.add_run(' -');fmt(p,9,align=WD_ALIGN_PARAGRAPH.CENTER,after=0)
doc.core_properties.title='진주 주거지의 서울 조기 도착 접근성 분석과 장거리 첫차 연계 개선'
doc.core_properties.author=IDENTITY['participant']
doc.core_properties.subject='공식 붙임2 항목과 표 구조에 따른 분석보고서'
doc.core_properties.comments=''
output=ROOT/'output/submission'/f"{IDENTITY['participant']}_분석보고서.docx"
doc.save(output)
print('Official form DOCX:',output)
