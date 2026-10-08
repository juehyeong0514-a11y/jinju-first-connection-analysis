"""Extract public workbook cells needed for reproducible rail transfer checks."""
import json
from datetime import time
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];RAW=ROOT/'data/raw/strengthening'
def minute(v):
    return v.hour*60+v.minute+v.second/60 if isinstance(v,time) else None
w=load_workbook(RAW/'korail_suin_bundang.xlsx',read_only=True,data_only=True)
records=[]
for sheet,days in [('평일상행',['DAY']),('휴일상행',['SAT','END'])]:
    rows=list(w[sheet].values)
    for i,row in enumerate(rows):
        if str(row[0]).strip() not in ['수서','선릉']:continue
        for j,train in enumerate(rows[3]):
            if not str(train).startswith('K'):continue
            for day in days:
                records.append(dict(line='수인분당',station=str(row[0]).strip(),day=day,train=train,direction='UP',arr=minute(row[j]),dep=minute(rows[i+1][j]),source_row=f'{sheet}!R{i+1}/R{i+2}C{j+1}'))
(RAW/'korail_suin_records.json').write_text(json.dumps(records,ensure_ascii=False,indent=2))
w=load_workbook(RAW/'korail_ktx_20261001.xlsx',read_only=True,data_only=True)
rows=list(w['경전선'].values); services=[]
for i,r in enumerate(rows):
    if len(r)>45 and r[24] in [204,382,'204','382']:
        services.append(dict(train=str(r[24]),jinju=minute(r[26]),suseo=minute(r[41]),gwangmyeong=minute(r[42]),seoul=minute(r[43]),operating_days=r[45],source_row=i+1))
(RAW/'korail_ktx_check.json').write_text(json.dumps(services,ensure_ascii=False,indent=2))
print('Extracted',len(records),'suburban station records; KTX checks',services)
